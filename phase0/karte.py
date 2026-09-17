# -*- coding: utf-8 -*-
"""カルテ_リスクマップ シートの読み取りと正規化（規則バージョン norm/v1）。

Phase 0 の受入は「既存 xlsx を投入 → DB → 再出力した xlsx の内容が入力と一致」。
一致の判定はバイナリ比較ではできない（作成日時・XML の順序・スタイルで必ず差が出る）。
何を一致と見なすかを規則として固定し、規則自体をテスト対象にする。
規則の全文は phase0/NORMALIZATION.md。ここはその実装。
"""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata

import openpyxl

NORM_VERSION = 'norm/v1'
KARTE_SHEET = 'カルテ_リスクマップ'
MASTER_SHEET = 'リスクマップマスタ'
AUTO_SHEETS = ('リスクマップ_AUTO', 'ヒートマップ_AUTO')

# 正準列名は build_risk_map.py（既存資産・無改変で流用）の入力仕様に合わせる。
CANONICAL = ['RiskItem', 'Big', 'Mid', 'Frame', 'Summary',
             'ProbBefore', 'ImpactBefore', 'Action', 'ProbAfter', 'ImpactAfter']

# 実テンプレートの列名は builder と食い違う。別名で吸収する。
# 同じ正準名へ 2 列が写る／未知の列がある／必須列が無い場合は全てエラーにする。
ALIASES = {
    'RiskItem': 'RiskItem',
    'Big': 'Big', 'BigCategory': 'Big',
    'Mid': 'Mid', 'MidCategory': 'Mid',
    'Frame': 'Frame', 'SmallFrame': 'Frame',
    'Summary': 'Summary',
    'ProbBefore': 'ProbBefore',
    'ImpactBefore': 'ImpactBefore',
    'Action': 'Action', 'ActionPlan': 'Action',
    'ProbAfter': 'ProbAfter',
    'ImpactAfter': 'ImpactAfter',
}

INT_COLS = ('ProbBefore', 'ImpactBefore', 'ProbAfter', 'ImpactAfter')
TEXT_COLS = ('RiskItem', 'Big', 'Mid', 'Frame', 'Summary', 'Action')
# 業務キー。DB 側の (domain, theme, measure, frame, summary) に対応する。
BUSINESS_KEY = ('RiskItem', 'Big', 'Mid', 'Frame', 'Summary')
FRAMES = ('管理可能性', '精度', 'スピード')

_WS_RUN = re.compile(r'[ \t]+')
_EXCEL_ERRORS = ('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NULL!', '#NUM!')


class KarteError(Exception):
    """入力が正規化規則に反している。黙って落とさず必ず送出する。"""


def norm_text(v, where):
    """文字列の正規化。NFC → 空白の統一 → 前後除去 → 内部の連続空白を 1 個へ。

    **NFKC は使わない。** NFKC は全角括弧「（）」を半角へ、全角英数字を半角へ
    畳んでしまい、台帳の値そのものを書き換える（実測: '人事・労務（Phase1）' が
    '人事・労務(Phase1)' になった）。往復では辻褄が合うが、DB に入る値が
    原本と変わるのは正規化ではなく改変なので採らない。
    合成済み・分解済みの揺れ（濁点等）だけを畳む NFC を使い、空白は下で明示的に扱う。

    大小文字は変換しない（日本語主体の台帳で誤変換の害が大きい）。
    NULL と空文字は同一視して '' を返す。
    """
    if v is None:
        return ''
    if isinstance(v, bool):
        raise KarteError(f'{where}: 真偽値は想定していない')
    if isinstance(v, (int, float)):
        # 数値セルに入った文字列項目。数値として書かれていても文字列として扱う。
        v = format_number(v, where)
    if not isinstance(v, str):
        raise KarteError(f'{where}: 想定外の型 {type(v).__name__}')
    if v in _EXCEL_ERRORS:
        raise KarteError(f'{where}: Excel のエラー値 {v}')
    s = unicodedata.normalize('NFC', v)
    s = s.replace(' ', ' ').replace('　', ' ')   # NBSP / 全角空白
    s = s.replace('\r\n', '\n').replace('\r', '\n')       # 改行は LF へ
    s = '\n'.join(_WS_RUN.sub(' ', line).strip() for line in s.split('\n'))
    return s.strip()


def format_number(v, where):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    raise KarteError(f'{where}: 非整数 {v!r} を文字列項目に使えない')


def norm_int_1_5(v, where):
    """1〜5 の整数。切り捨てない。非整数・範囲外・空はエラー。"""
    if v is None or (isinstance(v, str) and v.strip() == ''):
        raise KarteError(f'{where}: 必須の数値が空')
    if isinstance(v, bool):
        raise KarteError(f'{where}: 真偽値は想定していない')
    if isinstance(v, str):
        s = unicodedata.normalize('NFKC', v).strip()
        if s in _EXCEL_ERRORS:
            raise KarteError(f'{where}: Excel のエラー値 {s}')
        try:
            v = float(s)
        except ValueError:
            raise KarteError(f'{where}: 数値として読めない {v!r}')
    if isinstance(v, float):
        if not v.is_integer():
            raise KarteError(f'{where}: 非整数 {v!r}（切り捨てない）')
        v = int(v)
    if not isinstance(v, int):
        raise KarteError(f'{where}: 想定外の型 {type(v).__name__}')
    if not 1 <= v <= 5:
        raise KarteError(f'{where}: 1〜5 の範囲外 {v}')
    return v


def resolve_headers(raw_headers):
    """ヘッダ行を正準名へ写す。未知・不足・衝突・重複は全てエラー。"""
    mapping = {}          # 列インデックス -> 正準名
    seen = {}             # 正準名 -> 元の列名
    # 規則は「末尾の完全に空の列だけ無視」。位置を見ずに空ヘッダを飛ばすと、
    # 途中に空ヘッダの列があったとき、その列を丸ごと黙って取りこぼす。
    last_named = -1
    for idx, h in enumerate(raw_headers):
        if h is not None and str(h).strip() != '':
            last_named = idx
    for idx, h in enumerate(raw_headers):
        if h is None or str(h).strip() == '':
            if idx < last_named:
                raise KarteError(f'{idx + 1} 列目のヘッダが空（末尾以外の空ヘッダは許さない）')
            continue      # 末尾の空列だけ無視する
        name = unicodedata.normalize('NFKC', str(h)).strip()
        if name not in ALIASES:
            raise KarteError(f'未知の列: {name!r}')
        canon = ALIASES[name]
        if canon in seen:
            raise KarteError(f'列の衝突: {seen[canon]!r} と {name!r} が同じ {canon!r} に写る')
        seen[canon] = name
        mapping[idx] = canon
    missing = [c for c in CANONICAL if c not in seen]
    if missing:
        raise KarteError('必須の列が無い: ' + ', '.join(missing))
    return mapping


def read_karte(path):
    """カルテ_リスクマップ シートを正規化済みの行リストとして返す。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if KARTE_SHEET not in wb.sheetnames:
        raise KarteError(f'シートが無い: {KARTE_SHEET}')
    ws = wb[KARTE_SHEET]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise KarteError('カルテシートが空')
    mapping = resolve_headers(header)

    # 数式セルは受け付けない。data_only=True はキャッシュを読むだけで、
    # キャッシュが古くても検知できない（＝黙って古い値を通す）。
    formula_wb = openpyxl.load_workbook(path, data_only=False)
    fws = formula_wb[KARTE_SHEET]
    for row in fws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                raise KarteError(f'{cell.coordinate}: 数式セルは扱わない')

    out = []
    for rno, raw in enumerate(rows_iter, start=2):
        if raw is None or all(c is None or str(c).strip() == '' for c in raw):
            continue      # 全列空の行は行として数えない
        rec = {}
        for idx, canon in mapping.items():
            v = raw[idx] if idx < len(raw) else None
            where = f'行{rno} {canon}'
            rec[canon] = norm_int_1_5(v, where) if canon in INT_COLS else norm_text(v, where)
        for k in BUSINESS_KEY:
            if rec[k] == '':
                raise KarteError(f'行{rno}: 業務キー {k} が空')
        if rec['Frame'] not in FRAMES:
            raise KarteError(f'行{rno}: Frame は {"/".join(FRAMES)} のいずれか（{rec["Frame"]!r}）')
        out.append(rec)
    if not out:
        raise KarteError('カルテシートにデータ行が無い')

    keys = [tuple(r[k] for k in BUSINESS_KEY) for r in out]
    dup = {k for k in keys if keys.count(k) > 1}
    if dup:
        raise KarteError(f'業務キーの重複が {len(dup)} 件（台帳として誤り）')
    return out


def sort_key(rec):
    """ロケール差を避けるため、コードポイント順で並べる。"""
    return tuple(rec[k] for k in BUSINESS_KEY)


def serialize(rows):
    """正規化済みデータの直列化。JSON Lines・キー順固定・非 ASCII はそのまま。"""
    lines = []
    for rec in sorted(rows, key=sort_key):
        ordered = {k: rec[k] for k in CANONICAL}
        lines.append(json.dumps(ordered, ensure_ascii=False, sort_keys=True,
                                separators=(',', ':')))
    return '\n'.join(lines) + '\n'


def digest(rows):
    """直列化した UTF-8 バイト列の SHA-256（16 進小文字）。"""
    return hashlib.sha256(serialize(rows).encode('utf-8')).hexdigest()


def compare(left_rows, right_rows):
    """差分の一覧を返す。空リストなら「差分 0 件」。"""
    li = {sort_key(r): r for r in left_rows}
    ri = {sort_key(r): r for r in right_rows}
    diffs = []
    for k in sorted(set(li) - set(ri)):
        diffs.append(('入力にのみ存在', k, None, None))
    for k in sorted(set(ri) - set(li)):
        diffs.append(('出力にのみ存在', k, None, None))
    for k in sorted(set(li) & set(ri)):
        for col in CANONICAL:
            if li[k][col] != ri[k][col]:
                diffs.append((f'{col} が不一致', k, li[k][col], ri[k][col]))
    return diffs
